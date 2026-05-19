"""
sync_nkis.py — synchronizace dat z Národního kardiologického informačního systému (NKIS)

Stahuje xlsx datové souhrny ze stránek ÚZIS ČR, parsuje je
a ukládá jako JSON soubory do data/nkis/.

Spouští se přes GitHub Actions 1× měsíčně (viz .github/workflows/sync-nkis.yml)
nebo ručně:    python scripts/sync_nkis.py

Zavislosti:    pip install requests openpyxl
"""

import json
import sys
from pathlib import Path
from datetime import datetime
import requests
from openpyxl import load_workbook

# Adresář, kam ukládat JSON výstupy (relativně k umístění tohoto skriptu)
OUT_DIR = Path(__file__).parent.parent / "data" / "nkis"

# Definice datasetů — URL na xlsx souhrny ÚZIS a jak je interpretovat
# Poznámka: pokud ÚZIS přejmenuje soubor nebo změní strukturu, aktualizuj zde.
DATASETS = {
    "aim": {
        "label": "Akutní infarkt myokardu",
        "code": "I21-I22",
        "url": "https://datanzis.uzis.gov.cz/data/OIS-01-NKIS/OIS-01-14/Datovy-souhrn-OIS-01-14-akutni-infarkt-myokardu.xlsx",
        "sheet": "Hospitalizace_a_umrti",  # název listu v xlsx, který obsahuje data
        "year_col": "A",  # sloupec s rokem
        "value_col": "B",  # sloupec s počtem případů
        "header_row": 1,  # řádek s hlavičkami
    },
    "fs": {
        "label": "Fibrilace síní",
        "code": "I48",
        "url": "https://datanzis.uzis.gov.cz/data/OIS-01-NKIS/OIS-01-15/Datovy-souhrn-OIS-01-15-fibrilace-sini.xlsx",
        "sheet": "Lecene_osoby",
        "year_col": "A",
        "value_col": "B",
        "header_row": 1,
    },
    "hf": {
        "label": "Srdeční selhání",
        "code": "I50",
        "url": "https://datanzis.uzis.gov.cz/data/OIS-01-NKIS/OIS-01-16/Datovy-souhrn-OIS-01-16-srdecni-selhani.xlsx",
        "sheet": "Lecene_osoby",
        "year_col": "A",
        "value_col": "B",
        "header_row": 1,
    },
    "i35": {
        "label": "Aortální chlopeň (nerevmatická)",
        "code": "I35",
        "url": "https://datanzis.uzis.gov.cz/data/OIS-01-NKIS/OIS-01-17/Datovy-souhrn-OIS-01-17-aortalni-chlopen.xlsx",
        "sheet": "Lecene_osoby",
        "year_col": "A",
        "value_col": "B",
        "header_row": 1,
    },
    "i71": {
        "label": "Výduť aorty",
        "code": "I71",
        "url": "https://datanzis.uzis.gov.cz/data/OIS-01-NKIS/OIS-01-18/Datovy-souhrn-OIS-01-18-vydut-aorty.xlsx",
        "sheet": "Lecene_osoby",
        "year_col": "A",
        "value_col": "B",
        "header_row": 1,
    },
    "cmp": {
        "label": "Cévní mozková příhoda",
        "code": "I60-I64",
        "url": "https://datanzis.uzis.gov.cz/data/OIS-01-NKIS/OIS-01-19/Datovy-souhrn-OIS-01-19-cmp.xlsx",
        "sheet": "Hospitalizace_a_umrti",
        "year_col": "A",
        "value_col": "B",
        "header_row": 1,
    },
    "hyp": {
        "label": "Hypertenze (léčená)",
        "code": "I10",
        "url": "https://datanzis.uzis.gov.cz/data/OIS-01-NKIS/OIS-01-20/Datovy-souhrn-OIS-01-20-hypertenze.xlsx",
        "sheet": "Lecene_osoby",
        "year_col": "A",
        "value_col": "B",
        "header_row": 1,
    },
}


def download_xlsx(url: str, dest: Path) -> bool:
    """Stáhne xlsx ze zadané URL. Vrací True při úspěchu."""
    try:
        print(f"  Stahuji {url}...")
        response = requests.get(url, timeout=60, headers={
            "User-Agent": "Omnimedia-NKIS-Sync/1.0 (PR research)"
        })
        response.raise_for_status()
        dest.write_bytes(response.content)
        return True
    except Exception as e:
        print(f"  CHYBA při stahování: {e}", file=sys.stderr)
        return False


def parse_xlsx(xlsx_path: Path, cfg: dict) -> list:
    """Vyextrahuje časovou řadu (year, value) z xlsx souboru."""
    wb = load_workbook(xlsx_path, data_only=True)

    # Najdi správný list — pokud zadané jméno neexistuje, vezmi první
    sheet_name = cfg["sheet"]
    if sheet_name not in wb.sheetnames:
        print(f"  POZOR: list '{sheet_name}' nenalezen, používám '{wb.sheetnames[0]}'")
        sheet_name = wb.sheetnames[0]

    ws = wb[sheet_name]

    series = []
    header_row = cfg["header_row"]
    year_col = cfg["year_col"]
    value_col = cfg["value_col"]

    # Procházej řádky od header_row + 1 dál
    for row in range(header_row + 1, ws.max_row + 1):
        year_cell = ws[f"{year_col}{row}"].value
        value_cell = ws[f"{value_col}{row}"].value

        # Akceptujeme jen řádky, kde rok je integer mezi 2000 a 2030
        if isinstance(year_cell, (int, float)) and 2000 <= int(year_cell) <= 2030:
            if isinstance(value_cell, (int, float)):
                series.append({
                    "year": int(year_cell),
                    "value": int(value_cell)
                })

    return series


def compute_meta(series: list) -> dict:
    """Spočítá meta-informace pro frontend (trend, delta, peak)."""
    if not series or len(series) < 2:
        return {"trend": "unknown", "delta": 0, "peakYear": None}

    first = series[0]
    last = series[-1]
    delta_pct = round(((last["value"] - first["value"]) / first["value"]) * 100)

    peak = max(series, key=lambda r: r["value"])

    if delta_pct > 10:
        trend = "up"
    elif delta_pct < -10:
        trend = "down"
    else:
        trend = "plateau"

    return {
        "trend": trend,
        "delta": delta_pct,
        "peakYear": peak["year"]
    }


def sync_one(dataset_id: str, cfg: dict, tmp_dir: Path) -> bool:
    """Stáhne, naparsuje a uloží jeden dataset. Vrací True při úspěchu."""
    print(f"\n[{dataset_id}] {cfg['label']} ({cfg['code']})")

    xlsx_path = tmp_dir / f"{dataset_id}.xlsx"
    if not download_xlsx(cfg["url"], xlsx_path):
        return False

    try:
        series = parse_xlsx(xlsx_path, cfg)
        if not series:
            print(f"  CHYBA: žádná data v xlsx", file=sys.stderr)
            return False
    except Exception as e:
        print(f"  CHYBA při parsování: {e}", file=sys.stderr)
        return False

    meta = compute_meta(series)

    out = {
        "id": dataset_id,
        "label": cfg["label"],
        "code": cfg["code"],
        "source": "NRHZS",
        "source_type": "national",
        "updated": datetime.now().strftime("%Y-%m-%d"),
        "coverage": f"{series[0]['year']}–{series[-1]['year']}",
        "trend": meta["trend"],
        "delta": meta["delta"],
        "peakYear": meta["peakYear"],
        "data": series,
        "source_url": cfg["url"],
    }

    out_path = OUT_DIR / f"{dataset_id}.json"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(out, ensure_ascii=False, indent=2))

    print(f"  OK: {len(series)} datových bodů, delta {meta['delta']:+}%, "
          f"uloženo do {out_path.relative_to(OUT_DIR.parent.parent)}")
    return True


def main():
    print(f"NKIS sync — start v {datetime.now().isoformat()}")

    tmp_dir = Path("/tmp") / "nkis_sync"
    tmp_dir.mkdir(parents=True, exist_ok=True)

    results = {}
    for ds_id, cfg in DATASETS.items():
        results[ds_id] = sync_one(ds_id, cfg, tmp_dir)

    success_count = sum(1 for ok in results.values() if ok)
    total_count = len(results)

    print(f"\n=== NKIS sync — hotovo: {success_count}/{total_count} datasetů ===")

    if success_count == 0:
        print("\nVŠECHNY DATASETY SELHALY. Pravděpodobné příčiny:", file=sys.stderr)
        print("  1. ÚZIS přejmenoval xlsx soubory (zkontroluj URL v DATASETS)", file=sys.stderr)
        print("  2. ÚZIS server je dočasně nedostupný", file=sys.stderr)
        print("  3. Změnila se struktura xlsx (sheet name, sloupce)", file=sys.stderr)
        sys.exit(1)

    if success_count < total_count:
        failed = [k for k, v in results.items() if not v]
        print(f"\nČástečné selhání. Neúspěšné: {', '.join(failed)}")
        sys.exit(2)


if __name__ == "__main__":
    main()
